import asyncio
import io

from docx import Document
from pypdf import PdfReader


class OCRUnavailableError(RuntimeError):
    """The PDF renderer needed for vision transcription is unavailable."""


ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def decode_text(data: bytes) -> str:
    """UTF-8 first (with BOM support), then Windows-1251, common for older Ukrainian files."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def extract_text(data: bytes, ext: str) -> str:
    """Plain text of a .txt, .pdf or .docx file. Blocking: run it with asyncio.to_thread."""
    if ext == ".txt":
        return decode_text(data)
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n\n".join(p for p in parts if p.strip())
    raise ValueError(f"Unsupported extension: {ext}")


def _pdf_pages(data: bytes) -> list[tuple[str, bool]]:
    """Native text and whether a textless page needs visual extraction."""
    reader = PdfReader(io.BytesIO(data))
    result = []
    for page in reader.pages:
        text = page.extract_text() or ""
        result.append((text, not text.strip() and bool(page.images)))
    return result


def _render_page(data: bytes, index: int) -> list[bytes]:
    try:
        import pymupdf
    except ImportError as exc:
        raise OCRUnavailableError("Install PyMuPDF to render scanned PDFs") from exc
    with pymupdf.open(stream=data, filetype="pdf") as document:
        page = document[index]
        # Limit image size while keeping small print legible.
        scale = min(2.5, 2400 / max(page.rect.width, page.rect.height))
        # Two overlapping views give small text more visual tokens. Overlap
        # prevents lines at the split from being lost; the prompt merges them.
        rect = page.rect
        images = []
        for start, end in ((0, 0.55), (0.45, 1)):
            clip = pymupdf.Rect(rect.x0, rect.y0 + rect.height * start,
                                rect.x1, rect.y0 + rect.height * end)
            images.append(page.get_pixmap(matrix=pymupdf.Matrix(scale, scale),
                                          clip=clip, colorspace=pymupdf.csRGB,
                                          alpha=False).tobytes("png"))
        return images


async def extract_document(data: bytes, ext: str) -> str:
    """Native extraction for text documents; vision transcription for scans."""
    if ext != ".pdf":
        return await asyncio.to_thread(extract_text, data, ext)
    from . import llm

    pages = await asyncio.to_thread(_pdf_pages, data)
    texts = []
    for index, (text, needs_vision) in enumerate(pages):
        if needs_vision:
            images = await asyncio.to_thread(_render_page, data, index)
            text = await llm.transcribe_page(images)
        texts.append(text)
    return "\n\n".join(texts)
